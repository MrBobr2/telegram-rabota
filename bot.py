import asyncio
import os
import logging
import sqlite3
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import Command
from aiogram.types import ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO

# ==================== ЛОГИРОВАНИЕ ====================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s"
)
logger = logging.getLogger(__name__)

# ==================== ТОКЕН ====================
TOKEN = os.getenv("BOT_TOKEN", "")
if not TOKEN:
    logger.error("BOT_TOKEN не установлен! Добавьте переменную окружения BOT_TOKEN в Railway.")
    exit(1)

# ==================== БАЗА ДАННЫХ ====================
DB_PATH = "/data/tasks.db"

db_dir = os.path.dirname(DB_PATH)
if db_dir and not os.path.exists(db_dir):
    logger.warning(f"Папка {db_dir} не найдена. Пробую создать...")
    try:
        os.makedirs(db_dir, exist_ok=True)
    except Exception as e:
        logger.error(f"Не удалось создать папку {db_dir}: {e}. Проверьте, что Volume примонтирован.")

try:
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    cursor = conn.cursor()
    logger.info(f"Подключение к БД установлено: {DB_PATH}")
except Exception as e:
    logger.error(f"Не удалось открыть БД {DB_PATH}: {e}")
    exit(1)

try:
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            username TEXT,
            task_text TEXT,
            status TEXT,
            created_at TEXT,
            moved_to_waiting_at TEXT,
            completed_at TEXT,
            note TEXT
        )
    ''')
    conn.commit()
    logger.info("Таблица tasks готова.")
except Exception as e:
    logger.error(f"Ошибка при создании таблицы tasks: {e}")

try:
    cursor.execute("PRAGMA table_info(tasks)")
    columns = [row[1] for row in cursor.fetchall()]
    logger.info(f"Колонки в таблице tasks: {columns}")
    if "username" not in columns:
        logger.warning("В таблице tasks нет колонки username. Выполняю миграцию...")
        cursor.execute("ALTER TABLE tasks ADD COLUMN username TEXT")
        cursor.execute("UPDATE tasks SET username = 'Неизвестно' WHERE username IS NULL")
        conn.commit()
        logger.info("Миграция завершена: добавлена колонка username.")
except Exception as e:
    logger.error(f"Ошибка при проверке/миграции схемы: {e}")

# ==================== ФУНКЦИИ РАБОТЫ С БД ====================

def add_task(user_id, username, task_text):
    now = datetime.now().isoformat()
    try:
        cursor.execute(
            'INSERT INTO tasks (user_id, username, task_text, status, created_at) VALUES (?, ?, ?, ?, ?)',
            (user_id, username, task_text, 'in_work', now)
        )
        conn.commit()
        task_id = cursor.lastrowid
        logger.info(f"Задача добавлена: id={task_id}, user_id={user_id}, username={username}")
        return task_id
    except Exception as e:
        logger.error(f"Ошибка записи в БД (add_task): {e}")
        raise

def get_tasks(user_id, status):
    try:
        cursor.execute(
            'SELECT id, task_text, created_at, moved_to_waiting_at, completed_at, note, username '
            'FROM tasks WHERE user_id = ? AND status = ? ORDER BY id',
            (user_id, status)
        )
        rows = cursor.fetchall()
        tasks = []
        for row in rows:
            tasks.append({
                'id': row[0],
                'text': row[1],
                'created_at': datetime.fromisoformat(row[2]) if row[2] else None,
                'moved_to_waiting_at': datetime.fromisoformat(row[3]) if row[3] else None,
                'completed_at': datetime.fromisoformat(row[4]) if row[4] else None,
                'note': row[5],
                'username': row[6]
            })
        logger.debug(f"get_tasks(user_id={user_id}, status={status}) -> {len(tasks)} записей")
        return tasks
    except Exception as e:
        logger.error(f"Ошибка чтения из БД (get_tasks): {e}")
        return []

def get_task_by_id(task_id):
    try:
        cursor.execute(
            'SELECT id, task_text, created_at, moved_to_waiting_at, completed_at, note, status, username '
            'FROM tasks WHERE id = ?',
            (task_id,)
        )
        row = cursor.fetchone()
        if row:
            return {
                'id': row[0],
                'text': row[1],
                'created_at': datetime.fromisoformat(row[2]) if row[2] else None,
                'moved_to_waiting_at': datetime.fromisoformat(row[3]) if row[3] else None,
                'completed_at': datetime.fromisoformat(row[4]) if row[4] else None,
                'note': row[5],
                'status': row[6],
                'username': row[7]
            }
        return None
    except Exception as e:
        logger.error(f"Ошибка чтения из БД (get_task_by_id): {e}")
        return None

def update_task_status(task_id, status, extra_field=None, extra_value=None):
    try:
        if extra_field and extra_value is not None:
            cursor.execute(
                f'UPDATE tasks SET status = ?, {extra_field} = ? WHERE id = ?',
                (status, extra_value, task_id)
            )
        else:
            cursor.execute(
                'UPDATE tasks SET status = ? WHERE id = ?',
                (status, task_id)
            )
        conn.commit()
        logger.info(f"Статус задачи id={task_id} изменён на '{status}'")
    except Exception as e:
        logger.error(f"Ошибка обновления статуса (update_task_status): {e}")
        raise

def update_task_note(task_id, note):
    try:
        cursor.execute('UPDATE tasks SET note = ? WHERE id = ?', (note, task_id))
        conn.commit()
        logger.info(f"Пометка для задачи id={task_id} обновлена.")
    except Exception as e:
        logger.error(f"Ошибка обновления пометки (update_task_note): {e}")
        raise

def delete_task(task_id):
    try:
        cursor.execute('DELETE FROM tasks WHERE id = ?', (task_id,))
        conn.commit()
        logger.info(f"Задача id={task_id} удалена.")
    except Exception as e:
        logger.error(f"Ошибка удаления задачи (delete_task): {e}")
        raise

# ==================== ЭКСПОРТ В EXCEL ====================

def _fill_sheet(ws, tasks, include_completed=False):
    """Заполняет лист Excel данными задач (без столбца «Добавил в БД»)."""
    headers = ["№", "Задача", "Пометка", "Дата добавления", "Перенесена в ожидание"]
    if include_completed:
        headers.append("Дата выполнения")

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, task in enumerate(tasks, start=2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=task['text'])
        ws.cell(row=row_idx, column=3, value=task.get('note') or "")
        ws.cell(row=row_idx, column=4, value=format_date(task['created_at']))
        ws.cell(row=row_idx, column=5, value=format_date(task['moved_to_waiting_at']))

        if include_completed:
            ws.cell(row=row_idx, column=6, value=format_date(task['completed_at']))

    widths = [5, 50, 30, 20, 25]
    if include_completed:
        widths.append(20)

    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    if tasks:
        ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def create_excel_report(user_id):
    """Создаёт Excel-файл со всеми задачами пользователя. Возвращает BytesIO."""
    wb = Workbook()

    ws_work = wb.active
    ws_work.title = "В работе"
    _fill_sheet(ws_work, get_tasks(user_id, "in_work"), include_completed=False)

    ws_waiting = wb.create_sheet("В ожидании")
    _fill_sheet(ws_waiting, get_tasks(user_id, "waiting"), include_completed=False)

    ws_completed = wb.create_sheet("Выполненные")
    _fill_sheet(ws_completed, get_tasks(user_id, "completed"), include_completed=True)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

# ==================== ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ====================

class TaskStates(StatesGroup):
    waiting_for_note = State()

def get_main_menu():
    kb = ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="🔄 В работе"), KeyboardButton(text="⏸️ В ожидании")],
            [KeyboardButton(text="✅ Выполненные"), KeyboardButton(text="🗑️ Удалить выполненные")],
            [KeyboardButton(text="📊 Экспорт в Excel")]
        ],
        resize_keyboard=True,
        input_field_placeholder="Выберите действие"
    )
    return kb

def format_date(dt):
    if dt is None:
        return "Не указано"
    return dt.strftime("%d.%m.%Y %H:%M")

def get_username(message: types.Message) -> str:
    if message.from_user.first_name:
        return message.from_user.first_name
    if message.from_user.username:
        return f"@{message.from_user.username}"
    return f"User {message.from_user.id}"

# ==================== ПАГИНАЦИЯ ====================

PAGE_SIZE = 10

def build_paginated_keyboard(tasks, prefix, page=0, action_emoji="🗑️"):
    """Строит inline-клавиатуру с постраничным выводом."""
    total = len(tasks)
    total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
    page = max(0, min(page, total_pages - 1))

    start = page * PAGE_SIZE
    end = start + PAGE_SIZE
    page_tasks = tasks[start:end]

    keyboard = []
    for task in page_tasks:
        text = f"{action_emoji} {task['text'][:40]}"
        if len(task['text']) > 40:
            text += "..."
        keyboard.append([InlineKeyboardButton(text=text, callback_data=f"{prefix}_{task['id']}")])

    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️ Назад", callback_data=f"page_{prefix}_{page - 1}"))
    if page < total_pages - 1:
        nav_row.append(InlineKeyboardButton(text="Вперёд ➡️", callback_data=f"page_{prefix}_{page + 1}"))
    if nav_row:
        keyboard.append(nav_row)

    keyboard.append([InlineKeyboardButton(
        text=f"Стр. {page + 1} из {total_pages} (всего {total})",
        callback_data="noop"
    )])

    return InlineKeyboardMarkup(inline_keyboard=keyboard)

# ==================== БОТ И ДИСПЕТЧЕР ====================
bot = Bot(token=TOKEN)
storage = MemoryStorage()
dp = Dispatcher(storage=storage)

# ==================== ОБРАБОТЧИКИ ====================

@dp.message(Command("start"))
async def cmd_start(message: types.Message):
    user_name = message.from_user.first_name
    await message.answer(
        f"👋 Привет, {user_name}!\n"
        f"Я бот для управления задачами.\n\n"
        f"📝 Отправьте мне сообщение с задачей.\n"
        f"📌 Чтобы добавить пометку к заказу, используйте кнопку '📝 Пометка' в разделе задач.\n"
        f"📊 Для выгрузки всех задач в Excel используйте кнопку 'Экспорт в Excel'.\n\n"
        f"Все задачи сохраняются в базе данных и не пропадают при перезагрузке.",
        reply_markup=get_main_menu()
    )

@dp.message(Command("menu"))
async def cmd_menu(message: types.Message):
    await message.answer("Главное меню:", reply_markup=get_main_menu())

# ---------- Просмотр задач ----------

@dp.message(F.text == "🔄 В работе")
async def show_in_work(message: types.Message):
    user_id = message.from_user.id
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await message.answer("📭 Нет задач в работе.", reply_markup=get_main_menu())
        return

    response = "🔄 **Задачи в работе:**\n\n"
    for i, task in enumerate(tasks, 1):
        if task.get('note'):
            response += f"📌 **Пометка:** {task['note']}\n"
        response += f"{i}. {task['text']}\n"
        response += f"   📅 Добавлена: {format_date(task['created_at'])}\n\n"

    inline_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Выполнить", callback_data="show_complete_work")],
        [InlineKeyboardButton(text="⏸️ В ожидание", callback_data="show_waiting_transfer")],
        [InlineKeyboardButton(text="📝 Пометка", callback_data="show_add_note_work")]
    ])
    await message.answer(response, reply_markup=get_main_menu())
    await message.answer("Действия с задачами:", reply_markup=inline_kb)

@dp.message(F.text == "⏸️ В ожидании")
async def show_waiting_tasks(message: types.Message):
    user_id = message.from_user.id
    tasks = get_tasks(user_id, "waiting")
    if not tasks:
        await message.answer("📭 Нет задач в ожидании.", reply_markup=get_main_menu())
        return

    response = "⏸️ **Задачи в ожидании:**\n\n"
    for i, task in enumerate(tasks, 1):
        if task.get('note'):
            response += f"📌 **Пометка:** {task['note']}\n"
        response += f"{i}. {task['text']}\n"
        response += f"   📅 Добавлена: {format_date(task['created_at'])}\n"
        response += f"   ⏸️ Перенесена в ожидание: {format_date(task['moved_to_waiting_at'])}\n\n"

    inline_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Выполнить", callback_data="show_complete_waiting")],
        [InlineKeyboardButton(text="📝 Пометка", callback_data="show_add_note_waiting")]
    ])
    await message.answer(response, reply_markup=get_main_menu())
    await message.answer("Действия с задачами:", reply_markup=inline_kb)

@dp.message(F.text == "✅ Выполненные")
async def show_completed_tasks(message: types.Message):
    user_id = message.from_user.id
    tasks = get_tasks(user_id, "completed")
    if not tasks:
        await message.answer("📭 Нет выполненных задач.", reply_markup=get_main_menu())
        return

    response = "✅ **Выполненные задачи:**\n\n"
    for i, task in enumerate(tasks, 1):
        if task.get('note'):
            response += f"📌 **Пометка:** {task['note']}\n"
        response += f"{i}. {task['text']}\n"
        response += f"   📅 Добавлена: {format_date(task['created_at'])}\n"
        response += f"   ✅ Выполнена: {format_date(task['completed_at'])}\n\n"

    inline_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📝 Пометка", callback_data="show_add_note_completed")]
    ])
    await message.answer(response, reply_markup=get_main_menu())
    await message.answer("Действия с задачами:", reply_markup=inline_kb)

# ---------- Заглушка для noop ----------

@dp.callback_query(F.data == "noop")
async def noop(callback: types.CallbackQuery):
    await callback.answer()

# ---------- Пометки: показ списка ----------

@dp.callback_query(F.data == "show_add_note_work")
async def show_add_note_work(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await callback.answer("Нет задач в работе", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "addnote_work", page=0, action_emoji="📝")
    await callback.message.answer("Выберите задачу для добавления пометки:", reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data == "show_add_note_waiting")
async def show_add_note_waiting(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    tasks = get_tasks(user_id, "waiting")
    if not tasks:
        await callback.answer("Нет задач в ожидании", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "addnote_waiting", page=0, action_emoji="📝")
    await callback.message.answer("Выберите задачу для добавления пометки:", reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data == "show_add_note_completed")
async def show_add_note_completed(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    tasks = get_tasks(user_id, "completed")
    if not tasks:
        await callback.answer("Нет выполненных задач", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "addnote_completed", page=0, action_emoji="📝")
    await callback.message.answer("Выберите задачу для добавления пометки:", reply_markup=inline_kb)
    await callback.answer()

# ---------- Пометки: пагинация ----------

@dp.callback_query(F.data.startswith("page_addnote_work_"))
async def paginate_addnote_work(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[3])
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await callback.answer("Нет задач в работе", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "addnote_work", page=page, action_emoji="📝")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("page_addnote_waiting_"))
async def paginate_addnote_waiting(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[3])
    tasks = get_tasks(user_id, "waiting")
    if not tasks:
        await callback.answer("Нет задач в ожидании", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "addnote_waiting", page=page, action_emoji="📝")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("page_addnote_completed_"))
async def paginate_addnote_completed(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[3])
    tasks = get_tasks(user_id, "completed")
    if not tasks:
        await callback.answer("Нет выполненных задач", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "addnote_completed", page=page, action_emoji="📝")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

# ---------- Пометки: выбор задачи и ввод ----------

@dp.callback_query(F.data.startswith("addnote_"))
async def process_add_note_choice(callback: types.CallbackQuery, state: FSMContext):
    parts = callback.data.split("_")
    task_id = int(parts[2])

    await state.update_data(task_id=task_id)
    await state.set_state(TaskStates.waiting_for_note)
    await callback.message.answer(
        "✏️ Напишите пометку для задачи:",
        reply_markup=types.ReplyKeyboardRemove()
    )
    await callback.answer()

@dp.message(TaskStates.waiting_for_note)
async def process_add_note(message: types.Message, state: FSMContext):
    note_text = message.text.strip()
    if not note_text:
        await message.answer(
            "⚠️ Пометка не может быть пустой. Напишите текст пометки:",
            reply_markup=types.ReplyKeyboardRemove()
        )
        return

    state_data = await state.get_data()
    task_id = state_data.get("task_id")
    if not task_id:
        await state.clear()
        await message.answer("⚠️ Ошибка. Попробуйте снова.", reply_markup=get_main_menu())
        return

    task = get_task_by_id(task_id)
    if not task:
        await state.clear()
        await message.answer("⚠️ Задача не найдена.", reply_markup=get_main_menu())
        return

    update_task_note(task_id, note_text)
    await message.answer(
        f"✅ Пометка добавлена к задаче:\n"
        f"📌 Пометка: {note_text}\n"
        f"📝 Задача: {task['text']}",
        reply_markup=get_main_menu()
    )
    await state.clear()

# ---------- Выполнение: показ списка ----------

@dp.callback_query(F.data == "show_complete_work")
async def show_complete_work(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await callback.answer("Нет задач в работе", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "complete_work", page=0, action_emoji="✅")
    await callback.message.answer("Выберите задачу для выполнения:", reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data == "show_complete_waiting")
async def show_complete_waiting(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    tasks = get_tasks(user_id, "waiting")
    if not tasks:
        await callback.answer("Нет задач в ожидании", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "complete_waiting", page=0, action_emoji="✅")
    await callback.message.answer("Выберите задачу для выполнения:", reply_markup=inline_kb)
    await callback.answer()

# ---------- Выполнение: пагинация ----------

@dp.callback_query(F.data.startswith("page_complete_work_"))
async def paginate_complete_work(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[3])
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await callback.answer("Нет задач в работе", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "complete_work", page=page, action_emoji="✅")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("page_complete_waiting_"))
async def paginate_complete_waiting(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[3])
    tasks = get_tasks(user_id, "waiting")
    if not tasks:
        await callback.answer("Нет задач в ожидании", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "complete_waiting", page=page, action_emoji="✅")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

# ---------- Выполнение: обработка выбора ----------

@dp.callback_query(F.data.startswith("complete_work_"))
async def process_complete_work(callback: types.CallbackQuery):
    task_id = int(callback.data.split("_")[2])
    task = get_task_by_id(task_id)
    if not task or task['status'] != 'in_work':
        await callback.answer("⚠️ Задача не найдена", show_alert=True)
        return

    now = datetime.now()
    try:
        cursor.execute(
            'UPDATE tasks SET status = ?, completed_at = ?, note = NULL WHERE id = ?',
            ('completed', now.isoformat(), task_id)
        )
        conn.commit()
        logger.info(f"Задача id={task_id} выполнена (из 'в работе').")
    except Exception as e:
        logger.error(f"Ошибка выполнения задачи id={task_id}: {e}")
        await callback.answer("Ошибка при выполнении задачи", show_alert=True)
        return

    await callback.message.answer(
        f"✅ Задача '{task['text']}' выполнена!\n"
        f"📅 Дата выполнения: {format_date(now)}\n"
        f"Осталось задач в работе: {len(get_tasks(callback.from_user.id, 'in_work'))}",
        reply_markup=get_main_menu()
    )
    await callback.answer("Задача отмечена как выполненная")

@dp.callback_query(F.data.startswith("complete_waiting_"))
async def process_complete_waiting(callback: types.CallbackQuery):
    task_id = int(callback.data.split("_")[2])
    task = get_task_by_id(task_id)
    if not task or task['status'] != 'waiting':
        await callback.answer("⚠️ Задача не найдена", show_alert=True)
        return

    now = datetime.now()
    try:
        cursor.execute(
            'UPDATE tasks SET status = ?, completed_at = ?, note = NULL WHERE id = ?',
            ('completed', now.isoformat(), task_id)
        )
        conn.commit()
        logger.info(f"Задача id={task_id} выполнена (из 'в ожидании').")
    except Exception as e:
        logger.error(f"Ошибка выполнения задачи id={task_id}: {e}")
        await callback.answer("Ошибка при выполнении задачи", show_alert=True)
        return

    await callback.message.answer(
        f"✅ Задача '{task['text']}' выполнена!\n"
        f"📅 Дата выполнения: {format_date(now)}\n"
        f"Осталось задач в ожидании: {len(get_tasks(callback.from_user.id, 'waiting'))}",
        reply_markup=get_main_menu()
    )
    await callback.answer("Задача отмечена как выполненная")

# ---------- Перенос в ожидание: показ и пагинация ----------

@dp.callback_query(F.data == "show_waiting_transfer")
async def show_waiting_transfer(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await callback.answer("Нет задач в работе", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "waiting", page=0, action_emoji="⏸️")
    await callback.message.answer("Выберите задачу для переноса в ожидание:", reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("page_waiting_"))
async def paginate_waiting(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[2])
    tasks = get_tasks(user_id, "in_work")
    if not tasks:
        await callback.answer("Нет задач в работе", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "waiting", page=page, action_emoji="⏸️")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("waiting_"))
async def process_move_to_waiting(callback: types.CallbackQuery):
    task_id = int(callback.data.split("_")[1])
    task = get_task_by_id(task_id)
    if not task or task['status'] != 'in_work':
        await callback.answer("⚠️ Задача не найдена", show_alert=True)
        return

    now = datetime.now()
    update_task_status(task_id, 'waiting', 'moved_to_waiting_at', now.isoformat())

    response = f"⏸️ Задача '{task['text']}' перенесена в ожидание!\n"
    response += f"📅 Дата переноса: {format_date(now)}\n"
    if task.get('note'):
        response += f"📌 Пометка: {task['note']}\n"
    response += f"Задач в работе: {len(get_tasks(callback.from_user.id, 'in_work'))}\n"
    response += f"Задач в ожидании: {len(get_tasks(callback.from_user.id, 'waiting'))}"

    await callback.message.answer(response, reply_markup=get_main_menu())
    await callback.answer("Задача перенесена в ожидание")

# ---------- Удаление выполненных: показ и пагинация ----------

@dp.message(F.text == "🗑️ Удалить выполненные")
async def delete_completed_tasks(message: types.Message):
    user_id = message.from_user.id
    tasks = get_tasks(user_id, "completed")
    if not tasks:
        await message.answer("⚠️ Нет выполненных задач для удаления.", reply_markup=get_main_menu())
        return

    inline_kb = build_paginated_keyboard(tasks, "delete", page=0, action_emoji="🗑️")
    await message.answer("Выберите задачу для удаления:", reply_markup=inline_kb)

@dp.callback_query(F.data.startswith("page_delete_"))
async def paginate_delete(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    page = int(callback.data.split("_")[2])
    tasks = get_tasks(user_id, "completed")
    if not tasks:
        await callback.answer("Нет выполненных задач", show_alert=True)
        return
    inline_kb = build_paginated_keyboard(tasks, "delete", page=page, action_emoji="🗑️")
    await callback.message.edit_reply_markup(reply_markup=inline_kb)
    await callback.answer()

@dp.callback_query(F.data.startswith("delete_"))
async def process_delete_completed(callback: types.CallbackQuery):
    task_id = int(callback.data.split("_")[1])
    task = get_task_by_id(task_id)
    if not task or task['status'] != 'completed':
        await callback.answer("⚠️ Задача не найдена", show_alert=True)
        return

    delete_task(task_id)
    await callback.message.answer(
        f"🗑️ Задача '{task['text']}' удалена!\n"
        f"Осталось выполненных задач: {len(get_tasks(callback.from_user.id, 'completed'))}",
        reply_markup=get_main_menu()
    )
    await callback.answer("Задача удалена")

# ---------- Экспорт в Excel ----------

@dp.message(F.text == "📊 Экспорт в Excel")
async def export_to_excel(message: types.Message):
    user_id = message.from_user.id

    tasks_work = get_tasks(user_id, "in_work")
    tasks_waiting = get_tasks(user_id, "waiting")
    tasks_completed = get_tasks(user_id, "completed")

    if not tasks_work and not tasks_waiting and not tasks_completed:
        await message.answer("📭 У вас нет задач для экспорта.", reply_markup=get_main_menu())
        return

    await message.answer("⏳ Формирую Excel-файл...")

    try:
        buffer = create_excel_report(user_id)
    except Exception as e:
        logger.error(f"Ошибка создания Excel: {e}")
        await message.answer("⚠️ Не удалось создать файл. Попробуйте позже.", reply_markup=get_main_menu())
        return

    filename = f"tasks_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    document = types.BufferedInputFile(buffer.read(), filename=filename)

    total = len(tasks_work) + len(tasks_waiting) + len(tasks_completed)
    await message.answer_document(
        document=document,
        caption=(
            f"📊 Экспорт задач\n\n"
            f"🔄 В работе: {len(tasks_work)}\n"
            f"⏸️ В ожидании: {len(tasks_waiting)}\n"
            f"✅ Выполнено: {len(tasks_completed)}\n"
